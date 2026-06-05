"""
Active DMC Contact Report Agent — Windows Desktop App
Double-click to run (or build to .exe with PyInstaller).
"""

import io
import re
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

import anthropic
from docx import Document
from docx.shared import Pt, Inches

# ── Prompt ────────────────────────────────────────────────────────────────────
PROMPT_TEMPLATE = """You are the Active DMC Contact Report Agent. Your job is to transform meeting notes into a polished contact report following Active DMC's exact format.

The meeting data may come from any source — a transcript, a Read AI report, hand-written notes, an email summary, bullet points, or any other format. Work with whatever is provided.

Here is the meeting data you will be working with:

<meeting_data>
{meeting_data}
</meeting_data>

Your output must follow this EXACT structure — do not deviate:

Hi [First name of main client contact],

Sharing notes from our call today.

Present on behalf of [Client Company]:          [Client attendees, comma separated]
Present on behalf of Active DMC:                [Active DMC attendees, comma separated]
Date:                                           [Date, e.g. 3rd December 2025]
Time:                                           [Time] UAE time
Location:                                       [Teams / Zoom / In person]

Summary of Discussion and Next Steps

[Topic heading]:
•\t[Discussion point]
•\t[Discussion point]
\to\t[Sub-point if needed]

[Topic heading]:
•\t[Discussion point]
•\t[Discussion point]

Next Steps

For Active DMC:
•\t[Action item owned by Active DMC]
•\t[Action item owned by Active DMC]

For [Client Company]:
•\t[Action item owned by client]
•\t[Action item owned by client]

Let us know if we missed out on anything.

Rules — follow these exactly:
- Use "Hi [first name]," not the full name
- Opening line: "Sharing notes from our call today." (or "Capturing action points from our call today." if action-focused)
- The section heading is "Summary of Discussion and Next Steps" — one combined heading, NOT two separate headings
- Under the summary, group points by TOPIC. Each topic heading ends with a colon (:) and is on its own line
- Use • for main bullet points under each topic
- Use o (indented) only for genuine sub-points under a bullet
- The "Next Steps" sub-section is at the bottom, after all topics
- Next Steps are split as "For Active DMC:" and "For [Client Company]:" — use bullet points, NOT numbered lists
- Closing line is always: "Let us know if we missed out on anything."
- Keep tone professional, concise, and client-friendly
- Organize discussion points by topic, not chronologically
- Preserve all specific details: names, dates, deadlines, publications, clients, campaign details, numbers
- Carefully assign each action item to the correct owner (Active DMC or client)
- Use placeholders like [Date], [Time], [Location] only if the information is genuinely missing
- Always specify "UAE time" after the time
- Do not invent information not present in the meeting data

Before writing, think through:
1. Client first name and company name
2. All attendees from both sides
3. Meeting date, time, location
4. All topics discussed — group points under each topic
5. All action items — who owns each one

Then write the complete contact report following the exact format above.

<contact_report>
"""


# ── .docx builder ─────────────────────────────────────────────────────────────
def build_docx(report_text: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    for line in report_text.split("\n"):
        stripped = line.rstrip()

        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped in ("Summary of Discussion and Next Steps", "Next Steps"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
            continue

        if stripped.startswith("•"):
            content = stripped.lstrip("•").lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(content)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run(f"• {content}")
            continue

        if re.match(r"^o[\t ]", stripped):
            content = stripped[1:].lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet 2")
                p.add_run(content)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.add_run(f"○ {content}")
            continue

        if re.match(r"^\d+\.\t", stripped):
            p = doc.add_paragraph()
            run = p.add_run(stripped.replace("\t", "  "))
            run.bold = True
            continue

        doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Main window ───────────────────────────────────────────────────────────────
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Active DMC · Contact Report Agent")
        self.geometry("900x750")
        self.resizable(True, True)
        self.configure(bg="#f5f5f5")
        self._report_text = ""
        self._build_ui()

    def _build_ui(self):
        # ── Header ──
        hdr = tk.Frame(self, bg="#1a1a2e", pady=12)
        hdr.pack(fill="x")
        tk.Label(
            hdr,
            text="Active DMC  ·  Contact Report Agent",
            bg="#1a1a2e",
            fg="white",
            font=("Segoe UI", 14, "bold"),
        ).pack()

        body = tk.Frame(self, bg="#f5f5f5", padx=20, pady=16)
        body.pack(fill="both", expand=True)

        # ── API key ──
        tk.Label(body, text="Anthropic API Key", bg="#f5f5f5",
                 font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.api_key_var = tk.StringVar()
        key_entry = ttk.Entry(body, textvariable=self.api_key_var, show="•", width=70)
        key_entry.pack(fill="x", pady=(2, 12))

        # ── Transcript input ──
        tk.Label(body, text="Paste meeting transcript / notes below",
                 bg="#f5f5f5", font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.transcript_box = tk.Text(
            body, height=14, font=("Consolas", 10),
            wrap="word", relief="solid", bd=1,
        )
        self.transcript_box.pack(fill="both", expand=True, pady=(2, 10))

        # ── Buttons ──
        btn_row = tk.Frame(body, bg="#f5f5f5")
        btn_row.pack(fill="x", pady=(0, 10))

        self.generate_btn = tk.Button(
            btn_row,
            text="Generate Report",
            command=self._start_generation,
            bg="#1a1a2e", fg="white",
            font=("Segoe UI", 10, "bold"),
            padx=18, pady=6, relief="flat", cursor="hand2",
        )
        self.generate_btn.pack(side="left")

        self.save_btn = tk.Button(
            btn_row,
            text="Save as .docx",
            command=self._save_docx,
            bg="#2e7d32", fg="white",
            font=("Segoe UI", 10, "bold"),
            padx=18, pady=6, relief="flat", cursor="hand2",
            state="disabled",
        )
        self.save_btn.pack(side="left", padx=(10, 0))

        self.clear_btn = tk.Button(
            btn_row,
            text="Clear",
            command=self._clear,
            bg="#757575", fg="white",
            font=("Segoe UI", 10),
            padx=14, pady=6, relief="flat", cursor="hand2",
        )
        self.clear_btn.pack(side="left", padx=(10, 0))

        # ── Status ──
        self.status_var = tk.StringVar(value="Enter your API key and paste a transcript to begin.")
        tk.Label(body, textvariable=self.status_var, bg="#f5f5f5",
                 fg="#555", font=("Segoe UI", 9)).pack(anchor="w", pady=(0, 6))

        # ── Output ──
        tk.Label(body, text="Contact Report", bg="#f5f5f5",
                 font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.output_box = tk.Text(
            body, height=16, font=("Consolas", 10),
            wrap="word", relief="solid", bd=1,
            state="disabled",
        )
        self.output_box.pack(fill="both", expand=True, pady=(2, 0))

    # ── Generation ──────────────────────────────────────────────────────────
    def _start_generation(self):
        api_key = self.api_key_var.get().strip()
        transcript = self.transcript_box.get("1.0", "end").strip()

        if not api_key:
            messagebox.showwarning("Missing API Key", "Please enter your Anthropic API key.")
            return
        if not transcript:
            messagebox.showwarning("Missing Transcript", "Please paste a transcript first.")
            return

        self.generate_btn.config(state="disabled")
        self.save_btn.config(state="disabled")
        self._set_output("")
        self.status_var.set("Generating… please wait.")
        threading.Thread(target=self._generate, args=(api_key, transcript), daemon=True).start()

    def _generate(self, api_key: str, transcript: str):
        try:
            client = anthropic.Anthropic(api_key=api_key)
            prompt = PROMPT_TEMPLATE.format(meeting_data=transcript)
            full_text = ""

            with client.messages.stream(
                model="claude-opus-4-7",
                max_tokens=20000,
                thinking={"type": "adaptive"},
                messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
            ) as stream:
                for event in stream:
                    if (
                        hasattr(event, "type")
                        and event.type == "content_block_delta"
                        and hasattr(event, "delta")
                        and getattr(event.delta, "type", "") == "text_delta"
                    ):
                        full_text += event.delta.text
                        self.after(0, self._set_output, full_text)

            full_text = re.sub(r"\s*</contact_report>\s*$", "", full_text).strip()
            self._report_text = full_text
            self.after(0, self._set_output, full_text)
            self.after(0, self._on_done)

        except anthropic.AuthenticationError:
            self.after(0, self._on_error, "Invalid API key. Please check and try again.")
        except anthropic.APIConnectionError:
            self.after(0, self._on_error, "Could not reach Anthropic API. Check your internet connection.")
        except Exception as exc:
            self.after(0, self._on_error, f"Something went wrong: {exc}")

    def _on_done(self):
        self.generate_btn.config(state="normal")
        self.save_btn.config(state="normal")
        self.status_var.set("✅ Report ready — click 'Save as .docx' to download.")

    def _on_error(self, msg: str):
        self.generate_btn.config(state="normal")
        self.status_var.set(f"❌ {msg}")
        messagebox.showerror("Error", msg)

    # ── Output helpers ───────────────────────────────────────────────────────
    def _set_output(self, text: str):
        self.output_box.config(state="normal")
        self.output_box.delete("1.0", "end")
        self.output_box.insert("1.0", text)
        self.output_box.config(state="disabled")
        self.output_box.see("end")

    def _save_docx(self):
        if not self._report_text:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            initialfile="ADMC_ContactReport.docx",
        )
        if not path:
            return
        try:
            data = build_docx(self._report_text)
            with open(path, "wb") as f:
                f.write(data)
            self.status_var.set(f"✅ Saved: {path}")
        except Exception as exc:
            messagebox.showerror("Save failed", str(exc))

    def _clear(self):
        self.transcript_box.delete("1.0", "end")
        self._set_output("")
        self._report_text = ""
        self.save_btn.config(state="disabled")
        self.status_var.set("Enter your API key and paste a transcript to begin.")


if __name__ == "__main__":
    app = App()
    app.mainloop()
