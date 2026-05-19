#!/usr/bin/env python3
"""
Active DMC Contact Report Agent v1.0
=====================================

Run it, paste your transcript, get a polished contact report .docx.

Usage:
  python contact_report_agent.py

Requirements:
  pip install anthropic python-docx
"""

from __future__ import annotations

import io
import os
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import anthropic
except ImportError:
    print("ERROR: anthropic not installed. Run: pip install anthropic")
    sys.exit(1)

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

# ── Output folder ─────────────────────────────────────────────────────────────
HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "output" / "contact_reports"
OUT_DIR.mkdir(parents=True, exist_ok=True)

ADMC_FONT = "Trebuchet MS"

# ── Prompt ────────────────────────────────────────────────────────────────────
PROMPT_TEMPLATE = """You are the Active DMC Contact Report Agent. Your job is to transform meeting notes into a polished contact report following Active DMC's exact format.

The meeting data may come from any source — a transcript, a Read AI report, hand-written notes, an email summary, bullet points, a voice memo transcription, or any other format. Work with whatever is provided.

Here is the meeting data you will be working with:

<meeting_data>
{meeting_data}
</meeting_data>

Your output must follow this exact structure:

Hi [Client Name],

Thanks for your time – following our discussion please see the contact report below.

Present on behalf of [Client Company]:                       [Client attendees]

Present on behalf of Active DMC:                       [Active DMC attendees]

Date:                                                                [Date]

Time:                                                               [Time] UAE time

Location:                                                          [Teams/Zoom/In person]

Summary of Discussion

•\t[Topic title]
o\t[Summary point]
o\t[Summary point]

•\t[Topic title]
o\t[Summary point]
o\t[Summary point]

Next Steps

1.\tActive DMC
o\t[Action point owned by Active DMC]
o\t[Action point owned by Active DMC]

2.\t[Client Name/Client Company]
o\t[Action point owned by client]
o\t[Action point owned by client]

Please let us know if we missed anything or if you have any questions.

Important rules to follow:

- Follow the Active DMC contact report format exactly as shown above
- Use "•" (bullet points) for main discussion topics and "o" (sub-bullets) for summary points under each topic
- Use numbered lists (1., 2.) for Next Steps owners, then "o" for individual action items
- Keep the tone professional, concise, and client-friendly
- Extract and synthesize information from the meeting data; do not copy the transcript word-for-word
- Organize discussion points by topic, not chronologically
- Preserve all specific details accurately: names, dates, deadlines, publications, clients, campaign details, deliverables, numbers, etc.
- For action items, carefully determine whether Active DMC or the client is responsible and place them under the correct owner
- If attendees, date, time, or location information is missing from the meeting data, use placeholders like "[Client Name]", "[Date]", "[Time]", or "[Location - Teams/Zoom/In person]"
- Always specify "UAE time" after the time
- Do not invent or assume information that is not present in the meeting data
- Write this as a professional follow-up email, not a meeting transcript or verbatim summary

Before writing your final contact report, use the scratchpad below to:
1. Identify the client name and company
2. List attendees from both sides
3. Extract meeting logistics (date, time, location)
4. Identify main discussion topics and organize key points under each
5. Extract action items and assign them to the correct owner (Active DMC or client)

<scratchpad>
[Your analysis and organization of the meeting data goes here]
</scratchpad>

Now write the complete contact report following the exact format specified above. Your final output should be the complete, polished contact report ready to send to the client - do not include the scratchpad in your final answer.

<contact_report>
"""

# ── .docx builder ─────────────────────────────────────────────────────────────

def _set_run_font(run, *, size_pt=10, bold=False):
    run.font.name = ADMC_FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ADMC_FONT)
    rfonts.set(qn("w:hAnsi"), ADMC_FONT)
    rfonts.set(qn("w:cs"), ADMC_FONT)


def build_docx(report_text: str) -> bytes:
    doc = Document()

    # Force Trebuchet MS as the document default
    style = doc.styles["Normal"]
    style.font.name = ADMC_FONT
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ADMC_FONT)
    rfonts.set(qn("w:hAnsi"), ADMC_FONT)
    rfonts.set(qn("w:cs"), ADMC_FONT)

    for line in report_text.split("\n"):
        stripped = line.rstrip()

        if not stripped:
            doc.add_paragraph("")
            continue

        # Section headings
        if stripped in ("Summary of Discussion", "Next Steps"):
            p = doc.add_paragraph()
            _set_run_font(p.add_run(stripped), size_pt=11, bold=True)
            continue

        # Main bullets  •\t…
        if stripped.startswith("•"):
            content = stripped.lstrip("•").lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet")
                _set_run_font(p.add_run(content))
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                _set_run_font(p.add_run(f"• {content}"))
            continue

        # Sub-bullets  o\t…
        if re.match(r"^o[\t ]", stripped):
            content = stripped[1:].lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet 2")
                _set_run_font(p.add_run(content))
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                _set_run_font(p.add_run(f"○ {content}"))
            continue

        # Numbered owner lines  1.\t…
        if re.match(r"^\d+\.\t", stripped):
            p = doc.add_paragraph()
            _set_run_font(p.add_run(stripped.replace("\t", "  ")), bold=True)
            continue

        p = doc.add_paragraph()
        _set_run_font(p.add_run(stripped))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Claude streaming ──────────────────────────────────────────────────────────

def generate_report(transcript: str, api_key: str) -> str:
    client = anthropic.Anthropic(api_key=api_key)
    prompt = PROMPT_TEMPLATE.format(meeting_data=transcript)
    full_text = ""

    print("\n" + "─" * 70)
    print("CONTACT REPORT")
    print("─" * 70 + "\n")

    with client.messages.stream(
        model="claude-opus-4-7",
        max_tokens=20000,
        thinking={"type": "adaptive"},
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
    ) as stream:
        for event in stream:
            if (
                hasattr(event, "type")
                and event.type == "content_block_delta"
                and hasattr(event, "delta")
                and getattr(event.delta, "type", "") == "text_delta"
            ):
                chunk = event.delta.text
                full_text += chunk
                print(chunk, end="", flush=True)

    full_text = re.sub(r"\s*</contact_report>\s*$", "", full_text).strip()
    print("\n\n" + "─" * 70)
    return full_text


# ── Helpers ───────────────────────────────────────────────────────────────────

def banner():
    print("\n" + "=" * 70)
    print("  Active DMC  ·  Contact Report Agent  v1.0")
    print("=" * 70)
    print("  Model  : claude-opus-4-7  |  Adaptive thinking on")
    print(f"  Output : {OUT_DIR}")
    print("=" * 70 + "\n")


def get_api_key() -> str:
    env = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if env:
        print("  ✓ Using ANTHROPIC_API_KEY from environment\n")
        return env
    key = input("Anthropic API key: ").strip()
    if not key:
        print("API key required.")
        sys.exit(1)
    return key


def collect_transcript() -> str:
    print("\nPaste your transcript / meeting notes below.")
    print('Type  END  on its own line when finished.\n')
    lines = []
    while True:
        try:
            line = input()
        except EOFError:
            break
        if line.strip().upper() == "END":
            break
        lines.append(line)
    transcript = "\n".join(lines).strip()
    if not transcript:
        print("No transcript entered. Exiting.")
        sys.exit(1)
    return transcript


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    banner()

    api_key = get_api_key()
    transcript = collect_transcript()

    print("\n  Generating contact report…")

    try:
        report = generate_report(transcript, api_key)
    except anthropic.AuthenticationError:
        print("\n  ✗ Invalid API key. Please check and try again.")
        sys.exit(1)
    except anthropic.APIConnectionError:
        print("\n  ✗ Could not reach Anthropic API. Check your internet connection.")
        sys.exit(1)
    except Exception as e:
        print(f"\n  ✗ Error: {e}")
        sys.exit(1)

    # Save .docx
    ts = datetime.now().strftime("%Y-%m-%d_%H%M")
    fname = f"ADMC_ContactReport_{ts}.docx"
    out_path = OUT_DIR / fname

    try:
        out_path.write_bytes(build_docx(report))
        print(f"  ✓ Saved: {out_path}")
    except Exception as e:
        print(f"  ✗ Could not save .docx: {e}")
        txt_path = out_path.with_suffix(".txt")
        txt_path.write_text(report, encoding="utf-8")
        print(f"  ✓ Saved as .txt instead: {txt_path}")

    print("\n" + "=" * 70 + "\n")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\nInterrupted.")
        sys.exit(1)
