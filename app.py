"""
Active DMC Contact Report Agent — Web App
Run with: streamlit run app.py
"""

import io
import re
from datetime import datetime
import streamlit as st
import anthropic
from docx import Document
from docx.shared import Pt, Inches

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Active DMC · Contact Report Agent",
    page_icon="📋",
    layout="wide",
)

# ── Prompt template ───────────────────────────────────────────────────────────
PROMPT_TEMPLATE = """You are the Active DMC Contact Report Agent. Your job is to transform meeting notes into a polished contact report following Active DMC's exact format.

The meeting data may come from any source — a transcript, a Read AI report, hand-written notes, an email summary, bullet points, or any other format. Work with whatever is provided.

Here is the meeting data you will be working with:

<meeting_data>
{meeting_data}
</meeting_data>

Your output must follow this EXACT structure — do not deviate:

Hi [First name of main client contact],

Sharing notes from our call today.

Present on behalf of [Client Company]:          [Client attendees, comma separated]
Present on behalf of Active DMC:                [Active DMC attendees, comma separated]
Date:                                           [Date, e.g. 3rd December 2025]
Time:                                           [Time] UAE time
Location:                                       [Teams / Zoom / In person]

Summary of Discussion and Next Steps

[Topic heading]:
•\t[Discussion point]
•\t[Discussion point]
\to\t[Sub-point if needed]

[Topic heading]:
•\t[Discussion point]
•\t[Discussion point]

Next Steps

For Active DMC:
•\t[Action item owned by Active DMC]
•\t[Action item owned by Active DMC]

For [Client Company]:
•\t[Action item owned by client]
•\t[Action item owned by client]

Let us know if we missed out on anything.

Rules — follow these exactly:
- Use "Hi [first name]," not the full name
- Opening line: "Sharing notes from our call today." (or "Capturing action points from our call today." if the meeting was very action-focused)
- The section heading is "Summary of Discussion and Next Steps" — one combined heading, NOT two separate headings
- Under the summary, group points by TOPIC. Each topic heading ends with a colon (:) and is on its own line
- Use • for main bullet points under each topic
- Use o (indented) only for genuine sub-points under a bullet
- The "Next Steps" sub-section is at the bottom, after all topics
- Next Steps are split as "For Active DMC:" and "For [Client Company]:" — use bullet points, NOT numbered lists
- Closing line is always: "Let us know if we missed out on anything."
- Keep tone professional, concise, and client-friendly
- Organize discussion points by topic, not chronologically
- Preserve all specific details: names, dates, deadlines, publications, clients, campaign details, numbers
- Carefully assign each action item to the correct owner (Active DMC or client)
- Use placeholders like [Date], [Time], [Location] only if the information is genuinely missing
- Always specify "UAE time" after the time
- Do not invent information not present in the meeting data

Before writing, think through:
1. Client first name and company name
2. All attendees from both sides
3. Meeting date, time, location
4. All topics discussed — group points under each topic
5. All action items — who owns each one

Then write the complete contact report following the exact format above.

<contact_report>
"""


def generate_docx(report_text: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    for line in report_text.split("\n"):
        stripped = line.rstrip()

        if not stripped:
            doc.add_paragraph("")
            continue

        # Section headings
        if stripped in ("Summary of Discussion and Next Steps", "Next Steps"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
            continue

        # "For Active DMC:" / "For [Client]:" owner headings
        if re.match(r"^For .+:$", stripped):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            continue

        # Topic headings ending with colon (e.g. "Reporting:")
        if stripped.endswith(":") and not stripped.startswith("•") and not stripped.startswith("o"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            continue

        # Main bullets  •\t…
        if stripped.startswith("•"):
            content = stripped.lstrip("•").lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(content)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run(f"• {content}")
            continue

        # Sub-bullets  o\t…
        if re.match(r"^o[\t ]", stripped):
            content = stripped[1:].lstrip("\t").strip()
            try:
                p = doc.add_paragraph(style="List Bullet 2")
                p.add_run(content)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.add_run(f"○ {content}")
            continue

        doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def get_api_key() -> str:
    return st.session_state.get("api_key", "").strip()


def generate_report(transcript: str) -> str:
    client = anthropic.Anthropic(api_key=get_api_key())
    prompt = PROMPT_TEMPLATE.format(meeting_data=transcript)

    output_placeholder = st.empty()
    full_text = ""

    with client.messages.stream(
        model="claude-opus-4-7",
        max_tokens=20000,
        thinking={"type": "adaptive"},
        messages=[
            {
                "role": "user",
                "content": [{"type": "text", "text": prompt}],
            }
        ],
    ) as stream:
        for event in stream:
            # Only surface text deltas — skip thinking blocks
            if (
                hasattr(event, "type")
                and event.type == "content_block_delta"
                and hasattr(event, "delta")
                and getattr(event.delta, "type", "") == "text_delta"
            ):
                full_text += event.delta.text
                output_placeholder.markdown(f"```\n{full_text}\n```")

    # Strip the closing </contact_report> tag if Claude included it
    full_text = re.sub(r"\s*</contact_report>\s*$", "", full_text).strip()
    output_placeholder.empty()
    return full_text


# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## 📋 Active DMC")
    st.markdown("### Contact Report Agent")
    st.divider()

    st.markdown("**Anthropic API Key**")
    api_key_input = st.text_input(
        label="api_key_input",
        type="password",
        placeholder="sk-ant-api03-...",
        value=st.session_state.get("api_key", ""),
        label_visibility="collapsed",
    )
    if api_key_input:
        st.session_state["api_key"] = api_key_input
        st.success("Key saved for this session.", icon="✅")
    else:
        st.info("Enter your Anthropic API key above to get started.", icon="🔑")

    st.divider()
    st.caption("Model: claude-opus-4-7\nAdaptive thinking on")


# ── Main UI ───────────────────────────────────────────────────────────────────
st.title("Contact Report Agent")
st.markdown(
    "Paste your meeting notes below, then click **Generate Report**."
)

transcript = st.text_area(
    label="transcript",
    placeholder=(
        "Paste anything — transcript, Read AI report, hand-written notes, email summary, "
        "bullet points, voice memo, or any combination."
    ),
    height=340,
    label_visibility="collapsed",
)

col_btn, col_status = st.columns([1, 5])
with col_btn:
    generate_clicked = st.button(
        "Generate Report",
        type="primary",
        use_container_width=True,
        disabled=not get_api_key(),
    )
with col_status:
    if not get_api_key():
        st.caption("⬅ Enter your API key in the sidebar first.")
    elif st.session_state.get("report"):
        st.caption("✅ Report ready — copy from the box below or download.")

st.divider()

# ── Generate ──────────────────────────────────────────────────────────────────
if generate_clicked:
    if not transcript.strip():
        st.warning("Please paste a transcript before generating.")
    else:
        try:
            with st.spinner("Generating contact report…"):
                report = generate_report(transcript)
            if report:
                st.session_state["report"] = report
                st.session_state["report_time"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        except anthropic.AuthenticationError:
            st.error("Invalid API key. Please check the key in the sidebar and try again.")
        except anthropic.APIConnectionError:
            st.error("Could not reach the Anthropic API. Check your internet connection.")
        except Exception as e:
            st.error(f"Something went wrong: {e}")

# ── Output ────────────────────────────────────────────────────────────────────
if st.session_state.get("report"):
    report = st.session_state["report"]
    ts = st.session_state.get("report_time", "")

    st.subheader("Contact Report")
    if ts:
        st.caption(f"Generated {ts}")

    st.text_area(
        label="report_output",
        value=report,
        height=540,
        label_visibility="collapsed",
    )

    col_dl, col_clear, _ = st.columns([1, 1, 4])
    with col_dl:
        filename = f"ADMC_ContactReport_{datetime.now().strftime('%Y-%m-%d_%H%M')}.docx"
        st.download_button(
            label="⬇ Download .docx",
            data=generate_docx(report),
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
    with col_clear:
        if st.button("Clear", use_container_width=True):
            st.session_state.pop("report", None)
            st.session_state.pop("report_time", None)
            st.rerun()
